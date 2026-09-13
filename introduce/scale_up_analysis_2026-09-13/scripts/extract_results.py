#!/usr/bin/env python3
"""只读提取既有 xlsx 结果，保留单元格、公式缓存及去重出处。

运行：bundled-python introduce/scale_up_analysis_2026-09-13/scripts/extract_results.py
依赖：openpyxl（仅 load_workbook，不 save）；不执行训练、不修改源工作簿。
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import warnings
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import mean

import openpyxl


def serial(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def write_csv(path, records, fields=None):
    fields = fields or list(records[0])
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[3])
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "data")
    args = parser.parse_args()
    root, output = args.root.resolve(), args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    names = ["RankMixer-汇总-0909.xlsx", "副本RankMixer-汇总-0909.xlsx", "RankMixer-汇总.xlsx"]
    workbooks, observations, blank_slots, bad_diffs = [], [], [], []
    for name in names:
        path = root / "docs/experiments" / name
        source = str(path.relative_to(root))
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        # Extensions are not preserved by an openpyxl *save*. This script never saves a workbook.
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
            values = openpyxl.load_workbook(path, data_only=True)
            formulas = openpyxl.load_workbook(path, data_only=False)
        item = {"path": source, "sha256": before, "sheets": []}
        for sheet in values:
            fsheet = formulas[sheet.title]
            cells = []
            for row in fsheet:
                for cell in row:
                    cache = sheet[cell.coordinate].value
                    if cell.value is not None or cache is not None:
                        cells.append({"cell": cell.coordinate, "value": serial(cache),
                                      "formula": cell.value if cell.data_type == "f" else None,
                                      "data_type": cell.data_type, "number_format": cell.number_format})
            sheet_item = {"name": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column,
                          "merged_ranges": sorted(str(r) for r in sheet.merged_cells.ranges), "cells": cells}
            item["sheets"].append(sheet_item)
            header_rows = [r for r in range(1, sheet.max_row + 1)
                           if sheet.cell(r, 1).value == "测试数据日期"]
            for hr in header_rows:
                end = next((r - 3 for r in header_rows if r > hr), sheet.max_row)
                model_columns = [c for c in range(2, sheet.max_column + 1)
                                 if str(sheet.cell(hr, c).value).strip() == "auc"]
                for r in range(hr + 1, end + 1):
                    test = sheet.cell(r, 1).value
                    if not isinstance(test, datetime):
                        continue
                    test_date = test.date()
                    if test_date.month == 7:
                        chain_start = date(2026, 7, 1)
                    elif test_date.month == 8:
                        chain_start = date(2026, 8, 15)
                    else:
                        raise ValueError(f"Unmapped chain: {test_date}")
                    for c in model_columns:
                        model = sheet.cell(hr - 2, c).value
                        if not model:
                            raise ValueError(f"Missing model header: {source}:{hr - 2}:{c}")
                        acell, ccell = sheet.cell(r, c), sheet.cell(r, c + 1)
                        basecell = sheet.cell(r, 2)
                        locus = {"source": source, "sheet": sheet.title, "auc_cell": acell.coordinate,
                                 "copc_cell": ccell.coordinate, "base_auc_cell": basecell.coordinate,
                                 "test_date_cell": f"A{r}", "model_header_cell": sheet.cell(hr - 2, c).coordinate,
                                 "metadata_cell": sheet.cell(hr - 1, c).coordinate,
                                 "metadata_raw": serial(sheet.cell(hr - 1, c).value)}
                        if not numeric(acell.value):
                            blank_slots.append({"chain_start": str(chain_start), "model_id": model,
                                                "test_date": str(test_date), "status": "no_record", **locus})
                            continue
                        if not numeric(basecell.value):
                            raise ValueError(f"Observed AUC without same-day baseline: {locus}")
                        gap = round(acell.value - basecell.value, 12)
                        diffcell = sheet.cell(r, c + 2) if c != 2 else None
                        if diffcell:
                            locus["reported_gap_cell"] = diffcell.coordinate
                            locus["reported_gap_value"] = diffcell.value
                            locus["reported_gap_formula"] = fsheet[diffcell.coordinate].value if fsheet[diffcell.coordinate].data_type == "f" else None
                            if numeric(diffcell.value) and abs(diffcell.value - gap) > 1e-12:
                                bad_diffs.append({**locus, "recomputed_gap": gap})
                        observations.append({"chain_start": str(chain_start), "model_id": model,
                                             "train_date": str(test_date - timedelta(days=1)), "test_date": str(test_date),
                                             "training_day": (test_date - chain_start).days,
                                             "dense_start": "cold" if test_date == chain_start + timedelta(days=1) else "warm",
                                             "auc": acell.value, "copc": ccell.value if numeric(ccell.value) else None,
                                             "base_auc": basecell.value, "auc_gap_vs_base": gap,
                                             "auc_gap_bp": round(gap * 10000, 8), "is_baseline": c == 2,
                                             "status": "observed", "location": locus})
        values.close()
        formulas.close()
        assert hashlib.sha256(path.read_bytes()).hexdigest() == before, "Source workbook modified"
        workbooks.append(item)
    grouped = defaultdict(list)
    for observation in observations:
        grouped[(observation["chain_start"], observation["model_id"], observation["test_date"])].append(observation)
    conflicts, records = [], []
    compare_fields = ["auc", "copc", "base_auc"]
    for key, group in sorted(grouped.items()):
        first = group[0]
        for other in group[1:]:
            if any(first[k] != other[k] for k in compare_fields):
                conflicts.append({"key": key, "records": group})
        record = {k: v for k, v in first.items() if k != "location"}
        record.update({"source_count": len(group), "canonical_source": first["location"]["source"],
                       "canonical_sheet": first["location"]["sheet"],
                       "auc_cell": first["location"]["auc_cell"], "copc_cell": first["location"]["copc_cell"],
                       "base_auc_cell": first["location"]["base_auc_cell"],
                       "sources_json": json.dumps([x["location"] for x in group], ensure_ascii=False, separators=(",", ":"))})
        records.append(record)
    if conflicts:
        (output / "result_conflicts.json").write_text(json.dumps(conflicts, ensure_ascii=False, indent=2), encoding="utf-8")
        raise ValueError("Conflicting observed metrics; do not silently choose or average them")
    assert not bad_diffs, f"AUC diff audit failed: {bad_diffs}"
    write_csv(output / "experiment_results.csv", records)
    # A separate planned file avoids interpreting missing measurements as numeric observations.
    prereg = root / "introduce/rankmixer_ablation_20260814_results.csv"
    with prereg.open(encoding="utf-8-sig", newline="") as f:
        planned = list(csv.DictReader(f))
    assert all(not row["auc"] and row["status"] == "not_run" for row in planned)
    write_csv(output / "preregistered_not_run.csv", planned)

    # Fixed-date pairings, never compare each model's unmatched full-window mean.
    pairs = [
        ("2026-07-01", "rankmixer_v1", "bn_rankmixer_v1"),
        ("2026-07-01", "bn_rankmixer_v1", "bn_rankmixer_v2"),
        ("2026-07-01", "bn_rankmixer_v2", "bn_rankmixer_v3"),
        ("2026-07-01", "bn_rankmixer_v3", "bn_rankmixer_v4"),
        ("2026-07-01", "bn_rankmixer_v3", "bn_rankmixer_v5"),
        ("2026-08-15", "bn_rankmixer_v5", "bn_rankmixer_v6"),
        ("2026-08-15", "bn_rankmixer_v6", "bn_rankmixer_v6_e2"),
        ("2026-08-15", "bn_rankmixer_v6_e2", "bn_rankmixer_v6_e3"),
        ("2026-08-15", "bn_rankmixer_v6_e2", "bn_rankmixer_v6_e2_small"),
        ("2026-08-15", "bn_rankmixer_v6_e2_small", "bn_rankmixer_v6_e4"),
        ("2026-08-15", "bn_rankmixer_v6_e2_small", "bn_rankmixer_v6_e2_small_1"),
        ("2026-08-15", "bn_rankmixer_v6_e2_small_1", "bn_rankmixer_v6_e2_small_2"),
        ("2026-08-15", "bn_rankmixer_v6_e2_small", "bn_rankmixer_v6_e2_small_3"),
        ("2026-08-15", "bn_rankmixer_v6_e2", "bn_rankmixer_v6_e2_small_3"),
        ("2026-08-15", "bn_rankmixer_v6_e2_small_2", "bn_rankmixer_v6_e2_small_3"),
        ("2026-08-15", "cvr-senet-mature-rankmixer-v1", "cvr-senet-mature-rankmixer-v4"),
        ("2026-08-15", "cvr-senet-mature-rankmixer-v1", "cvr-senet-mature-rankmixer-v5"),
        ("2026-08-15", "cvr-senet-mature-rankmixer-v4", "cvr-senet-mature-rankmixer-v5"),
        ("2026-08-15", "bn_rankmixer_v6", "bn_rankmixer_v7"),
        ("2026-08-15", "bn_rankmixer_v6", "bn_rankmixer_v8"),
        ("2026-08-15", "bn_rankmixer_v6", "bn_rankmixer_v9"),
        ("2026-08-15", "bn_rankmixer_v6", "unimixer_v1"),
    ]
    pair_records = []
    for chain, left, right in pairs:
        l = {r["test_date"]: r for r in records if r["chain_start"] == chain and r["model_id"] == left}
        r = {r["test_date"]: r for r in records if r["chain_start"] == chain and r["model_id"] == right}
        dates = sorted(l.keys() & r.keys())
        deltas = [round(r[d]["auc"] - l[d]["auc"], 12) for d in dates]
        pair_records.append({"chain_start": chain, "left_model": left, "right_model": right,
                             "dates": ";".join(dates), "n_dates": len(dates),
                             "first_day_delta": deltas[0], "mean_delta": round(mean(deltas), 12),
                             "warm_mean_delta": round(mean(deltas[1:]), 12) if len(deltas) > 1 else None,
                             "min_delta": min(deltas), "max_delta": max(deltas),
                             "positive_days": sum(x > 0 for x in deltas),
                             "daily_deltas": ";".join(f"{x:.6f}" for x in deltas)})
    write_csv(output / "paired_comparisons.csv", pair_records)
    summaries = []
    by_model = defaultdict(list)
    for r in records:
        by_model[(r["chain_start"], r["model_id"])].append(r)
    for (chain, model), group in sorted(by_model.items()):
        group.sort(key=lambda x: x["test_date"])
        summaries.append({"chain_start": chain, "model_id": model, "n_dates": len(group),
                          "first_test_date": group[0]["test_date"], "last_test_date": group[-1]["test_date"],
                          "first_auc": group[0]["auc"], "first_gap": group[0]["auc_gap_vs_base"],
                          "last_auc": group[-1]["auc"], "last_gap": group[-1]["auc_gap_vs_base"],
                          "mean_auc": round(mean(r["auc"] for r in group), 12),
                          "mean_gap": round(mean(r["auc_gap_vs_base"] for r in group), 12),
                          "positive_days_vs_base": sum(r["auc_gap_vs_base"] > 0 for r in group),
                          "mean_abs_copc_error": round(mean(abs(r["copc"] - 1) for r in group if r["copc"] is not None), 12),
                          "canonical_auc_cells": ";".join(r["auc_cell"] for r in group)})
    write_csv(output / "model_result_summary.csv", summaries)
    source_keys = defaultdict(set)
    for observation in observations:
        source_keys[observation["location"]["source"]].add((observation["chain_start"], observation["model_id"], observation["test_date"]))
    logs = []
    for path in sorted((root / "docs/logs").glob("*.log")):
        lines = path.read_text(errors="replace").splitlines()
        metrics = [i + 1 for i, line in enumerate(lines) if re.search(r"auc\s*[=:]\s*[0-9]|copc\s*[=:]\s*[0-9]", line, re.I)]
        errors = [{"line": i + 1, "text": line[:300]} for i, line in enumerate(lines)
                  if re.search(r"(?:ModuleNotFoundError|ImportError|RuntimeError|ValueError|TypeError):|some bg process failed", line)]
        logs.append({"path": str(path.relative_to(root)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                     "line_count": len(lines), "metric_pattern_matching_line_numbers": metrics,
                     "explicit_errors": errors,
                     "note": "仅检索结果/显式错误线索；缺少最终测试指标不等价于零AUC或全部运行失败。"})
    audit = {"workbooks": [{"path": w["path"], "sha256": w["sha256"],
                            "sheet_shapes": [{k: s[k] for k in ("name", "max_row", "max_column")} for s in w["sheets"]]} for w in workbooks],
             "source_observation_count": len(observations), "deduplicated_observation_count": len(records),
             "duplicate_observation_count": len(observations) - len(records),
             "baseline_observation_count": sum(r["is_baseline"] for r in records),
             "candidate_observation_count": sum(not r["is_baseline"] for r in records),
             "candidate_unique_model_labels": len({r["model_id"] for r in records if not r["is_baseline"]}),
             "counts_by_source": dict(Counter(o["location"]["source"] for o in observations)),
             "unique_observations_by_source": {k: len(v) for k, v in source_keys.items()},
             "copy_matches_primary_all_nonempty_cell_records": workbooks[0]["sheets"] == workbooks[1]["sheets"],
             "older_workbook_unique_observations_not_in_primary": len(source_keys[workbooks[2]["path"]] - source_keys[workbooks[0]["path"]]),
             "counts_by_chain": dict(Counter(r["chain_start"] for r in records)),
             "workbook_metric_conflict_count": len(conflicts), "invalid_reported_diff_count": len(bad_diffs),
             "blank_auc_slot_count_with_duplicates": len(blank_slots), "preregistered_not_run_count": len(planned),
             "date_protocol_source": "docs/overview/background.md:27-50",
             "date_protocol_note": "train_date/training_day/dense_start按文档协议从测试日期推导，非逐run日志验证；warm为同方案逐日热启动。",
             "nonempty_cell_storage_note": "workbook_cells.json保留所有非空单元格和公式/缓存，未列出的单元格为空；blank_auc_slots保存日期×方案空记录。",
             "statistical_note": "多个日期是相关续训/测试链，非独立seed；均值按共同测试日等权，不是合并样本总体AUC。",
             "logs": logs}
    (output / "workbook_cells.json").write_text(json.dumps({"audit": audit, "workbooks": workbooks, "blank_auc_slots": blank_slots}, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "extraction_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
